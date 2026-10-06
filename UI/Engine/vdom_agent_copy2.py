import os
import base64
import json
import yaml
import time
from datetime import datetime
from typing import Dict, TypedDict, Any, List, Literal
from dotenv import load_dotenv
from pydantic import BaseModel, Field

# 🔴 DEBUG TOGGLE: Set to False later to reduce terminal noise
VERBOSE_DEBUG_MODE = True

try:
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.drawing.image import Image as xlImage
except ImportError:
    print("❌ Missing required libraries. Please run: pip install pandas openpyxl Pillow")
    exit(1)

try:
    import cv2
    import numpy as np
except ImportError:
    print("❌ Missing Computer Vision libraries. Please run: pip install opencv-python numpy")
    exit(1)

# --- Load Environment Variables ---
load_dotenv()
os.environ["AZURE_OPENAI_API_KEY"] = os.getenv("AZURE_API_KEY", "")
os.environ["AZURE_OPENAI_ENDPOINT"] = os.getenv("AZURE_ENDPOINT", "")
os.environ["OPENAI_API_VERSION"] = os.getenv("AZURE_API_VERSION", "2024-06-01")
AZURE_DEPLOYMENT_NAME = os.getenv("AZURE_DEPLOYMENT_NAME", "gpt-4o")

from langchain_openai import AzureChatOpenAI
from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.graph import StateGraph, START, END
from langchain_community.callbacks.manager import get_openai_callback

# --- Pydantic Schema for Strict Outputs ---
class VerificationResult(BaseModel):
    status: Literal["PASS", "FAIL"] = Field(
        description=(
            "Evaluate the UI state. Output 'PASS' ONLY if the visual changes inside the red bounding box "
            "exactly match the expected business task without unintended side effects. "
            "Output 'FAIL' if the expected change is missing, incorrect, displays an error state, or affects the wrong data."
        )
    )
    description: str = Field(
        description=(
            "Provide a comprehensive, evidence-based QA observation. You MUST include: "
            "1) A reference to the specific visual evidence (e.g., 'Inside the red bounding box...'), "
            "2) The exact business data or component affected (use domain vocabulary like 'SKU', 'Grid', 'Filter'), and "
            "3) A definitive statement on whether the observed state matches the expected intent. "
            "NEVER use vague phrasing like 'it worked' or 'the UI updated'."
        )
    )

# --- Dynamic Agent Memory Initialization ---
AGENT_MEMORY_FILE = ""
agent_memory = {}
REPORT_DIR = ""
DIFF_DIR = ""

def init_agent_memory(workspace, kb_filename):
    global AGENT_MEMORY_FILE, agent_memory
    base_name = os.path.basename(kb_filename).replace(".yaml", "")
    
    memory_filename = f"agent_memory_{base_name}.yaml"
    AGENT_MEMORY_FILE = os.path.join(workspace, memory_filename) if workspace else memory_filename
    
    if not os.path.exists(AGENT_MEMORY_FILE):
        with open(AGENT_MEMORY_FILE, 'w', encoding='utf-8') as f:
            yaml.dump({"learned_rules": {}, "cached_chunks": {}, "global_rules": [], "global_domain_context": ""}, f)
            
    with open(AGENT_MEMORY_FILE, 'r', encoding='utf-8') as f:
        agent_memory = yaml.safe_load(f) or {"learned_rules": {}, "cached_chunks": {}, "global_rules": [], "global_domain_context": ""}
        
    if "cached_chunks" not in agent_memory: agent_memory["cached_chunks"] = {}
    if "global_rules" not in agent_memory: agent_memory["global_rules"] = []
    if "global_domain_context" not in agent_memory: agent_memory["global_domain_context"] = ""

def save_agent_memory():
    global AGENT_MEMORY_FILE, agent_memory
    with open(AGENT_MEMORY_FILE, 'w', encoding='utf-8') as f:
        yaml.dump(agent_memory, f, sort_keys=False)

def encode_image(image_path: str) -> str:
    if not image_path or not os.path.exists(image_path): return None
    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode('utf-8')

# 🔴 COMPUTER VISION DIFF ENGINE
def create_cv_diff_image(before_path, after_path, step_id):
    if not before_path or not after_path or not os.path.exists(before_path) or not os.path.exists(after_path):
        return after_path 
        
    out_path = os.path.join(DIFF_DIR, f"diff_{step_id}_{int(time.time())}.png")
    
    img1 = cv2.imread(before_path)
    img2 = cv2.imread(after_path)
    if img1 is None or img2 is None: return after_path
    
    g1 = cv2.GaussianBlur(cv2.cvtColor(img1, cv2.COLOR_BGR2GRAY), (11, 11), 0)
    g2 = cv2.GaussianBlur(cv2.cvtColor(img2, cv2.COLOR_BGR2GRAY), (11, 11), 0)
    
    diff = cv2.absdiff(g1, g2)
    _, thresh = cv2.threshold(diff, 25, 255, cv2.THRESH_BINARY)
    
    kernel = np.ones((5,5), np.uint8)
    thresh = cv2.dilate(thresh, kernel, iterations=3)
    
    contours, _ = cv2.findContours(thresh, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    
    result = img2.copy()
    drawn = False
    for c in contours:
        x, y, w, h = cv2.boundingRect(c)
        if w > 20 and h > 20: 
            cv2.rectangle(result, (x, y), (x+w, y+h), (0, 0, 255), 4)
            drawn = True
            
    if drawn:
        cv2.imwrite(out_path, result)
        if VERBOSE_DEBUG_MODE:
            print(f"      [CV Engine] Red-Box Diff Image Generated: {out_path}")
        return out_path
    return after_path

# --- 1. Define the LangGraph State ---
class AgentState(TypedDict):
    chunk_id: str
    steps: List[dict]
    chunk_signature: str
    start_image_path: str
    end_image_path: str
    code_chunk_text: str
    enriched_images: dict
    ui_graph: dict              
    verification_plan: list     
    confidence_score: float
    confusing_step_id: str
    loop_count: int
    is_cached: bool
    verification_status: str
    bug_description: str
    micro_debug_result: str
    domain_context: str

def segment_steps(steps: List[dict]) -> List[List[dict]]:
    if not steps: return []
    if len(steps) <= 3: return [steps] 
    
    print("  🧠 [Semantic Segmenter] Analyzing code to build semantic chunks...")
    llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.0, model_kwargs={"response_format": {"type": "json_object"}})
    steps_text = "\n".join([f"StepID [{s['step_id']}]: {s['raw_code']}" for s in steps])
    
    system_prompt = SystemMessage(content=(
        "You are an AI Test Architect. Group sequential automation steps into semantic chunks. "
        "A chunk represents a single logical user goal (e.g., 'Applying a filter'). "
        "Chunks should ideally be 4 to 8 steps long. Break chunks AFTER state-committing actions (clicks/navigates). "
        'Output JSON: { "chunks": [ ["1_1", "1_2"], ["1_3"] ] }'
    ))
    try:
        response = llm.invoke([system_prompt, HumanMessage(content=steps_text)])
        raw_json = response.content.strip().replace("```json", "").replace("```", "").strip()
        parsed = json.loads(raw_json)
        
        step_map = {str(s['step_id']).strip(): s for s in steps}
        final_chunks = []
        for id_list in parsed.get("chunks", []):
            chunk = []
            for i in id_list:
                clean_i = str(i).replace("Step", "").replace("ID", "").replace("[", "").replace("]", "").strip()
                if clean_i in step_map: chunk.append(step_map[clean_i])
            if chunk: final_chunks.append(chunk)
        return final_chunks
    except Exception as e:
        return [steps[i:i + 5] for i in range(0, len(steps), 5)]

def domain_analyzer_node(state: AgentState):
    if state.get('domain_context') or state['is_cached']:
        return state
        
    print(f"  [Node: Domain Analyzer] Extracting Business & Domain Context...")
    llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.2)
    b64_start = encode_image(state["start_image_path"])
    if not b64_start:
        state['domain_context'] = "Unknown Domain"
        return state
        
    system_prompt = SystemMessage(content=(
        "You are an expert Enterprise Data QA Engineer. Look at this application dashboard. "
        "Deduce its core business purpose. Identify the major components (e.g., Grids, Charts, Sidebar Filters) "
        "and extract specific domain vocabulary (e.g., 'MLB', 'Out of Stock', 'Forecast', 'Demand Planning'). "
        "Write a concise, 3-4 sentence paragraph describing what this app does and what data it displays."
    ))
    
    user_prompt = HumanMessage(content=[
        {"type": "text", "text": "Analyze this dashboard and extract the business context:"},
        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64_start}", "detail": "high"}}
    ])
    try:
        response = llm.invoke([system_prompt, user_prompt])
        domain_context = response.content.strip()
        state['domain_context'] = domain_context
        agent_memory["global_domain_context"] = domain_context
        save_agent_memory()
        print(f"    └── Extracted Context: {domain_context[:100]}...")
    except Exception as e:
        print(f"    └── ⚠️ Domain Analysis Error: {e}")
        state['domain_context'] = "Failed to extract domain context."
    return state

def predictor_node(state: AgentState):
    if state['is_cached']:
        print(f"  [Node: Predictor] ⚡ Cache Hit! Skipping V-DOM Generation.")
        return state
        
    print(f"  [Node: Predictor] Building Virtual DOM & Verification Plan (Loop: {state['loop_count']})...")
    llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.1, model_kwargs={"response_format": {"type": "json_object"}})
    
    content_payload = [
        {"type": "text", "text": f"Current V-DOM State:\n{json.dumps(state.get('ui_graph', {}), indent=2)}"},
        {"type": "text", "text": f"Code Chunk to simulate:\n{state['code_chunk_text']}"}
    ]
    
    b64_start = encode_image(state["start_image_path"])
    if b64_start:
        content_payload.append({"type": "text", "text": "Starting UI State:"})
        content_payload.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64_start}", "detail": "high"}})
        
    if state["enriched_images"]:
        for s_id, path in state["enriched_images"].items():
            b64_intent = encode_image(path)
            if b64_intent:
                content_payload.append({"type": "text", "text": f"Crosshair for confusing Step {s_id}:"})
                content_payload.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64_intent}", "detail": "high"}})
    
    learned_context = "\n".join([f"- GLOBAL RULE: {rule}" for rule in agent_memory.get("global_rules", [])])
    for step in state["steps"]:
        rule = agent_memory["learned_rules"].get(step['raw_code'])
        if rule: learned_context += f"\n- Rule for `{step['raw_code']}`: {rule}"
        human_assertions = step.get('human_assertions')
        if human_assertions and isinstance(human_assertions, list):
            for assertion in human_assertions:
                learned_context += f"\n- MANDATORY HUMAN ASSERTION for `{step['raw_code']}`: {assertion}"
    
    system_prompt = SystemMessage(content=(
        "You are a Senior Data QA Architect. Your job is to translate mechanical code into BUSINESS logic verifications. "
        f"CRITICAL DOMAIN CONTEXT: {state.get('domain_context', 'None identified')}\n\n"
        "RULE 1: NEVER output mechanical verifications like 'Verify checkbox is checked' or 'Verify sidebar is open'. "
        "RULE 2: Your verification plan MUST explain the DATA and BUSINESS IMPACT. If a filter is clicked, explain WHAT data should change in the grids or charts. "
        "RULE 3: Use the domain vocabulary (e.g., MLB, Out of Stock, Forecasting, SKU, Sales).\n\n"
        "RULE 4: If you receive a human assertion, pay special attention and ensure you verify the human assertion."
        f"HUMAN OVERRIDES & GUIDELINES: {learned_context}\n\n"
        "Output JSON Schema:\n"
        "{\n"
        '  "chunk_business_intent": "E.g., The user is applying a hierarchy filter for the Team attribute to dynamically update the Out of Stock alerts grid.",\n'
        '  "ui_graph": {"Updated_VDOM_Here": {}},\n'
        '  "confidence_score": 0.0 to 1.0,\n'
        '  "confusing_step_id": null,\n'
        '  "verification_plan": [\n'
        '    {"step_id_checkpoint": "1_3", "what_to_verify": "Verify the \'Team\' hierarchy filter is actively applied to the dataset and highlighted in the UI."},\n'
        '    {"step_id_checkpoint": "END", "what_to_verify": "Verify the Alerts Summary grid has successfully re-rendered to display only inventory data relevant to the selected Team."}\n'
        '  ]\n'
        "}"
    ))
    
    try:
        response = llm.invoke([system_prompt, HumanMessage(content=content_payload)])
        raw_json = response.content.strip().replace("```json", "").replace("```", "").strip()
        result = json.loads(raw_json)
        
        state['ui_graph'] = result.get("ui_graph", state['ui_graph'])
        state['confidence_score'] = result.get("confidence_score", 1.0)
        state['confusing_step_id'] = result.get("confusing_step_id")
        state['verification_plan'] = result.get("verification_plan", [])
        return state
    except Exception as e:
        state['confidence_score'] = 1.0 
        state['verification_plan'] = [{"step_id_checkpoint": "END", "what_to_verify": "Verify the business data updated successfully."}]
        return state

def enricher_node(state: AgentState):
    target_id = str(state['confusing_step_id']).replace("Step", "").replace("ID", "").replace("[", "").replace("]", "").strip()
    target_step = next((s for s in state['steps'] if str(s['step_id']).strip() == target_id), None)
    if target_step and target_step.get("baseline_images", {}).get("intent"):
        state['enriched_images'][target_id] = target_step["baseline_images"]["intent"]
    else:
        state['confidence_score'] = 1.0
    state['loop_count'] += 1
    return state

def verifier_node(state: AgentState):
    print(f"  [Node: Verifier] Executing {len(state['verification_plan'])} Atomic CV-Checkpoints...")
    
    # 🔴 Use Structured Output with Pydantic
    llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.0)
    structured_llm = llm.with_structured_output(VerificationResult)
    
    all_verdicts = []
    chunk_failed = False
    fail_reasons = []
    image_map = {str(s['step_id']): {"before": s.get("baseline_images", {}).get("before"), "after": s.get("baseline_images", {}).get("after")} for s in state['steps']}
    
    for task in state['verification_plan']:
        step_val = str(task.get("step_id_checkpoint"))
        
        if step_val == "END":
            before_path = state["steps"][0].get("baseline_images", {}).get("before")
            after_path = state["end_image_path"]
        else:
            before_path = image_map.get(step_val, {}).get("before")
            after_path = image_map.get(step_val, {}).get("after")
            
        if not after_path: continue
        diff_path = create_cv_diff_image(before_path, after_path, step_val)
        
        print(f"    └── Checking Frame [{step_val}]: {task.get('what_to_verify')}")
        b64_img = encode_image(diff_path)
        system_prompt = SystemMessage(content=(
            "You are a strict QA Inspector. You receive ONE image and ONE specific task. "
            f"BUSINESS CONTEXT: {state.get('domain_context', 'None identified')}\n\n"
            "IMPORTANT: Look closely at the RED BOX drawn on the image. This highlights exactly what changed in the UI. "
            "When describing the outcome, use the business context vocabulary (e.g., name the specific metrics, products, or filters affected). "
            "Does the visual state inside or around the RED BOX match the task exactly?"
        ))
        user_prompt = HumanMessage(content=[
            {"type": "text", "text": f"Task: {task.get('what_to_verify')}"},
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64_img}", "detail": "high"}}
        ])
        
        # 🔴 Smart Retry Logic for the Schema Output
        max_retries = 3
        status = "FAIL"
        description = "Unknown Error"
        
        for attempt in range(max_retries):
            try:
                # Ask the LLM for structured output
                result = structured_llm.invoke([system_prompt, user_prompt])
                
                # Check if the result is actually instantiated properly
                if hasattr(result, 'status') and hasattr(result, 'description'):
                    status = result.status
                    description = result.description
                    break  # Success! Exit the retry loop
                else:
                    raise ValueError("LLM returned an empty or malformed Pydantic object.")
                    
            except Exception as e:
                if attempt < max_retries - 1:
                    print(f"      ⚠️ Format error. Retrying ({attempt+1}/{max_retries})...")
                    time.sleep(2) # Give it a slightly longer pause before retrying
                else:
                    # If we exhaust all retries, do NOT crash. Force a FAIL state gracefully.
                    print(f"      ❌ Final format error: Defaulting to FAIL. Details: {str(e)[:100]}")
                    status = "FAIL"
                    description = f"System Error: AI failed to produce valid JSON validation format after {max_retries} attempts."
        
        all_verdicts.append(f"[{status}] {task.get('what_to_verify')} - {description}")
        
        if status == "FAIL":
            chunk_failed = True
            fail_reasons.append(description)
            
    state['verification_status'] = "FAIL" if chunk_failed else "PASS"
    state['bug_description'] = "\n".join(all_verdicts)
    return state

def micro_debug_node(state: AgentState):
    print(f"  [Node: Micro-Debug] Chunk failed. Replaying timeline to isolate bug...")
    llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.1, model_kwargs={"response_format": {"type": "json_object"}})
    
    content_payload = [{"type": "text", "text": "The chunk execution failed. Here is the frame-by-frame replay."}]
    
    for step in state['steps']:
        content_payload.append({"type": "text", "text": f"\nExecuting Step {step['step_id']}: `{step['raw_code']}`"})
        intent_b64 = encode_image(step.get("baseline_images", {}).get("intent"))
        if intent_b64:
            content_payload.append({"type": "text", "text": "Element Clicked:"})
            content_payload.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{intent_b64}", "detail": "high"}})
        after_b64 = encode_image(step.get("baseline_images", {}).get("after"))
        if after_b64:
            content_payload.append({"type": "text", "text": "Resulting UI State:"})
            content_payload.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{after_b64}", "detail": "high"}})
            
    system_prompt = SystemMessage(content=(
        "You are an investigative QA Agent reviewing a supposedly failed test. "
        "Find the exact step where the UI failed to respond properly. "
        "CRITICAL RULE: If the code executed perfectly and the UI behaved correctly, the previous agent hallucinated the failure. "
        "If it is a false positive, output 'failed_step_id': 'NONE'.\n"
        "Output JSON:\n"
        "{\n"
        '  "failed_step_id": "step_id or NONE",\n'
        '  "root_cause": "Detailed explanation."\n'
        "}"
    ))
    
    try:
        response = llm.invoke([system_prompt, HumanMessage(content=content_payload)])
        raw_json = response.content.strip().replace("```json", "").replace("```", "").strip()
        result = json.loads(raw_json)
        state['micro_debug_result'] = f"Failed at Step {result.get('failed_step_id')}: {result.get('root_cause')}"
        print(f"    └── Isolated Bug: {state['micro_debug_result']}")
    except Exception:
        state['micro_debug_result'] = "Could not isolate bug via micro-debug."
    return state

# --- 4. Define Routing Logic ---
def route_prediction(state: AgentState):
    if state['confidence_score'] < 0.85 and state.get('confusing_step_id') and state['loop_count'] < 2:
        return "enricher"
    
    if not state['is_cached'] and state['confidence_score'] >= 0.85:
        agent_memory["cached_chunks"][state['chunk_signature']] = {
            "ui_graph": state.get('ui_graph', {}),
            "verification_plan": state.get('verification_plan', [])
        }
        save_agent_memory()
        
    return "verifier"

def route_verification(state: AgentState):
    if state['verification_status'] == "FAIL":
        return "micro_debug"
    return END

def build_agent_graph():
    workflow = StateGraph(AgentState)
    workflow.add_node("domain_analyzer", domain_analyzer_node)
    workflow.add_node("predictor", predictor_node)
    workflow.add_node("enricher", enricher_node)
    workflow.add_node("verifier", verifier_node)
    workflow.add_node("micro_debug", micro_debug_node)
    
    workflow.add_edge(START, "domain_analyzer")
    workflow.add_edge("domain_analyzer", "predictor")
    workflow.add_conditional_edges("predictor", route_prediction)
    workflow.add_edge("enricher", "predictor")
    workflow.add_conditional_edges("verifier", route_verification)
    workflow.add_edge("micro_debug", END)
    
    return workflow.compile()



def get_dynamic_test_metadata(text_corpus: str, fallback_section: str = "General"):
    """
    Uses LLM to dynamically determine BOTH the Feature Area and a high-level Test Scenario summary.
    """
    if not text_corpus.strip():
        return fallback_section, "Unnamed Test Scenario"

    try:
        # Force JSON response to get two structured fields back
        llm = AzureChatOpenAI(azure_deployment=AZURE_DEPLOYMENT_NAME, temperature=0.0, model_kwargs={"response_format": {"type": "json_object"}})
        
        system_prompt = SystemMessage(content=(
            "You are an expert QA Lead. Based on the provided test verification steps and observed results, extract two things:\n"
            "1. 'feature_area': The primary UI component being tested (concise 2-4 words, e.g., 'Hierarchy Filter', 'Alerts Summary').\n"
            "2. 'test_scenario': A high-level summary of the test's purpose (short, 1 concise sentence, e.g., 'Verify user can apply the hierarchy filter to update the grid').\n"
            "Output strictly in JSON format:\n"
            '{"feature_area": "...", "test_scenario": "..."}'
        ))
        
        user_prompt = HumanMessage(content=f"Test Context:\n\n{text_corpus}")
        response = llm.invoke([system_prompt, user_prompt])
        
        raw_json = response.content.strip().replace("```json", "").replace("```", "").strip()
        data = json.loads(raw_json)
        
        feature_area = data.get("feature_area", fallback_section).title().replace("'", "").replace('"', '')
        test_scenario = data.get("test_scenario", "Unnamed Test Scenario")
        
        return feature_area, test_scenario
    except Exception as e:
        print(f"      ⚠️ Metadata extraction fallback: {e}")
        return fallback_section, "Unnamed Test Scenario"


# --- 6. Orchestrator ---
def main():
    global REPORT_DIR, DIFF_DIR
    
    print("\n" + "="*70)
    print("🤖 NAIVE AGENT V3: DOMAIN-AWARE V-DOM & ATOMIC VERIFICATION")
    print("="*70)
    
    workspace = os.getenv("AUTO_WORKSPACE")
    if not workspace:
        workspace = input("Enter Workspace Directory:\n> ").strip()
    
    project_name = os.path.basename(workspace.strip("/\\")) if workspace else "default_project"
    
    REPORT_DIR = os.path.join("UI", "outputs", project_name, "Naive_Agent_Reports")
    DIFF_DIR = os.path.join(REPORT_DIR, "CV_Diffs")
    os.makedirs(DIFF_DIR, exist_ok=True)
    
    kb_input = os.getenv("AUTO_KB")
    if not kb_input:
        kb_input = input("Enter target Knowledge Base YAML:\n> ").strip()
        
    if not kb_input.endswith(".yaml"): 
        kb_input += ".yaml"
    
    kb_path = os.path.join(workspace, kb_input) if workspace else kb_input
    if not os.path.exists(kb_path):
        print(f"\n❌ Error: '{kb_path}' not found.")
        return
        
    init_agent_memory(workspace, kb_path)
    
    with open(kb_path, "r", encoding="utf-8") as f: 
        kb = yaml.safe_load(f)
        
    agent_graph = build_agent_graph()
    
    RUN_TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.basename(kb_path).replace(".yaml", "")
    excel_report_path = os.path.join(REPORT_DIR, f"{base_name}_Report_{RUN_TIMESTAMP}.xlsx")
    
    excel_rows = []
    total_chunks = 0
    total_bugs = 0
    ongoing_ui_graph = {}
    total_cost = 0.0
    total_tokens = 0
    
    for section in kb.get("sections", []):
        section_name = section.get('section_name', 'Unknown')
        print(f"\n📁 Processing Section: {section_name}")
        
        valid_steps = [s for s in section.get("steps", []) if s.get("baseline_images", {}).get("before")]
        chunks = segment_steps(valid_steps)
        
        for chunk_idx, chunk_steps in enumerate(chunks):
            chunk_id = f"{section_name}_Chunk_{chunk_idx+1}"
            code_text = "\n".join([f"Step {s['step_id']}: {s['raw_code']}" for s in chunk_steps])
            chunk_signature = "|".join([s['raw_code'] for s in chunk_steps])
            
            print(f"\n  📦 Running {chunk_id} ({len(chunk_steps)} Actions)")
            
            is_cached = False
            verification_plan = []
            if chunk_signature in agent_memory.get("cached_chunks", {}):
                is_cached = True
                cached_data = agent_memory["cached_chunks"][chunk_signature]
                ongoing_ui_graph = cached_data.get("ui_graph", ongoing_ui_graph)
                verification_plan = cached_data.get("verification_plan", [])
            
            initial_state = {
                "chunk_id": chunk_id,
                "steps": chunk_steps,
                "chunk_signature": chunk_signature,
                "start_image_path": chunk_steps[0]["baseline_images"]["before"],
                "end_image_path": chunk_steps[-1]["baseline_images"]["after"],
                "code_chunk_text": code_text,
                "enriched_images": {},
                "ui_graph": ongoing_ui_graph, 
                "verification_plan": verification_plan,
                "confidence_score": 1.0,
                "confusing_step_id": None,
                "loop_count": 0,
                "is_cached": is_cached,
                "verification_status": "",
                "bug_description": "",
                "micro_debug_result": "",
                "domain_context": agent_memory.get("global_domain_context", "")
            }
            
            chunk_cost = 0.0
            with get_openai_callback() as cb:
                final_state = agent_graph.invoke(initial_state)
                chunk_cost = cb.total_cost
                total_cost += chunk_cost
                total_tokens += cb.total_tokens
                
            ongoing_ui_graph = final_state.get('ui_graph', ongoing_ui_graph)
            final_pass_fail = "PASS" if final_state["verification_status"] == "PASS" else "FAIL"
            
            # Extract final images for this chunk to embed in the report
            chunk_start_img = chunk_steps[0].get("baseline_images", {}).get("before")
            chunk_intent_img = chunk_steps[-1].get("baseline_images", {}).get("intent")
            chunk_end_img = chunk_steps[-1].get("baseline_images", {}).get("after")

            if final_state["verification_status"] == "FAIL":
                if "Failed at Step NONE" in final_state.get('micro_debug_result', ''):
                    print("\n   ✨ [Auto-Resolve] Micro-Debug confirmed Agent 2 hallucinated. Auto-passing chunk.")
                    final_pass_fail = "PASS"
                    # 🔴 Update state so the Excel report reflects the override
                    final_state["bug_description"] = "[AUTO-RESOLVED FALSE POSITIVE]\n" + final_state.get("bug_description", "")
                else:
                    print("\n" + "!"*60)
                    print(f"🚨 BUG DETECTED IN CHUNK")
                    print(f"   Micro-Debug Isolation: {final_state['micro_debug_result']}")
                    
                    if os.getenv("AUTO_TRIAGE") == "true":
                        print("   🤖 [Auto-Triage] Logging bug automatically for UI Dashboard.")
                        total_bugs += 1
                    else:
                        # Manual Terminal Mode - Open images for review
                        try:
                            start_img = os.path.abspath(final_state["start_image_path"])
                            end_img = os.path.abspath(final_state["end_image_path"])
                            print(f"   📸 Opening images for review...")
                            os.startfile(start_img)
                            time.sleep(0.3)
                            os.startfile(end_img)
                        except Exception as e: 
                            print(f"   ⚠️ Could not auto-open images: {e}")
                        
                        while True:
                            choice = input("\n   👉 Is this a REAL bug? (y = Confirmed, n = False Positive, s = Skip): ").strip().lower()
                            if choice in ['y', 'n', 's']: break
                            
                        if choice == 'y':
                            print("   └── 🐞 Bug Confirmed! Logging.")
                            total_bugs += 1
                        elif choice == 'n':
                            print("\n   [Rule Builder]")
                            scope = input("   Apply this rule to:\n   (1) This specific chunk only.\n   (2) Globally across the entire app.\n   > ").strip()
                            correction = input("   What SHOULD the Agent have expected? \n   > ")
                            
                            if scope == '2':
                                agent_memory["global_rules"].append(correction)
                            else:
                                agent_memory["learned_rules"][chunk_steps[-1]['raw_code']] = correction
                                
                            agent_memory["cached_chunks"].pop(chunk_signature, None)
                            save_agent_memory()
                            final_pass_fail = "PASS"
                            # 🔴 Update state so the Excel report reflects the manual override
                            final_state["bug_description"] = "[HUMAN OVERRIDE: FALSE POSITIVE]\n" + final_state.get("bug_description", "")
                            
            total_chunks += 1
            
            observed_combined = final_state.get("bug_description", "")
            if final_state.get("micro_debug_result"):
                observed_combined += f"\n\nMicro-Debug:\n{final_state.get('micro_debug_result')}"
            expected_text = "\n".join([f"- {task['what_to_verify']}" for task in final_state.get("verification_plan", [])])
            
            # 🔴 PRE-EXCEL LOGIC VALIDATION 🔴
            bug_desc = final_state.get("bug_description", "")
            pass_count = bug_desc.count("[PASS]")
            fail_count = bug_desc.count("[FAIL]")
            total_steps = pass_count + fail_count
            
            if total_steps > 0:
                if fail_count == total_steps and final_pass_fail == "PASS":
                    print("   🚨 [Validation] Contradiction caught: All steps failed. Forcing final status to FAIL.")
                    final_pass_fail = "FAIL"
                    observed_combined += "\n\n[SYSTEM OVERRIDE: Final status forced to FAIL because all individual steps failed.]"
                elif pass_count == total_steps and final_pass_fail == "FAIL":
                    print("   🚨 [Validation] Contradiction caught: All steps passed. Forcing final status to PASS.")
                    final_pass_fail = "PASS"
                    observed_combined += "\n\n[SYSTEM OVERRIDE: Final status forced to PASS because all individual steps passed.]"



             # --- Dynamic Feature Area & Test Scenario Extraction ---
            combined_context = f"Expected Result:\n{expected_text}\n\nObserved Result:\n{observed_combined}\n\nActions:\n{code_text}"
            detected_feature_area, business_intent = get_dynamic_test_metadata(combined_context, fallback_section=section_name)



            # -------------------------------------------------------------------
            # ✨ ADDED HERE: Synthesize Specific Pass/Fail Reason
            # -------------------------------------------------------------------
            if final_pass_fail == "PASS":
                if "[AUTO-RESOLVED FALSE POSITIVE]" in final_state.get("bug_description", ""):
                    verdict_reason = "Auto-Resolved: Micro-debug verified all UI actions succeeded and confirmed previous failure was a false positive."
                elif "[HUMAN OVERRIDE: FALSE POSITIVE]" in final_state.get("bug_description", ""):
                    verdict_reason = "Human Override: Manually verified and passed with custom domain rule applied."
                else:
                    verdict_reason = "Passed: All expected UI changes, filter updates, and grid renders matched the verification plan."
            else:
                # Isolate the primary failure reason from micro-debug or failed checkpoints
                micro_res = final_state.get("micro_debug_result", "")
                if micro_res and "Failed at Step NONE" not in micro_res:
                    verdict_reason = f"Failed: {micro_res}"
                else:
                    # Extract the failure descriptions from individual tasks
                    failed_tasks = [
                        line.split(" - ", 1)[-1] 
                        for line in final_state.get("bug_description", "").split("\n") 
                        if "[FAIL]" in line
                    ]
                    if failed_tasks:
                        verdict_reason = f"Failed: {' | '.join(failed_tasks[:2])}"
                    else:
                        verdict_reason = "Failed: UI state did not reflect the expected business outcome."

            # Store the raw paths here, using the new descriptive names
            excel_rows.append({
                "Test Scenario": f"{business_intent.title()} ({chunk_id})",
                "Feature Area": detected_feature_area,  
                "User Actions": code_text,
                "Expected Result": expected_text,
                "Observed Result": observed_combined,
                "Before Image": chunk_start_img,
                "Intent Image": chunk_intent_img,
                "After Image": chunk_end_img,
                "Pass/Fail": final_pass_fail,
                "Pass/Fail Reason": verdict_reason,
                "Cost ($)": round(chunk_cost, 4)
            })



    # --- Generate Styled Excel Report with Embedded Images ---
    if excel_rows:
        df = pd.DataFrame(excel_rows)
        
        # 1. Add Grand Total Row
        total_row = pd.DataFrame([{
            "Test Scenario": "GRAND TOTAL",
            "Feature Area": "",
            "User Actions": "",
            "Expected Result": "",
            "Observed Result": "",
            "Before Image": "",
            "Intent Image": "",
            "After Image": "",
            "Pass/Fail": "",
            "Pass/Fail Reason": "",
            "Cost ($)": round(total_cost, 4)
        }])
        df = pd.concat([df, total_row], ignore_index=True)
        
        # 2. Write to Excel with openpyxl engine
        with pd.ExcelWriter(excel_report_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='QA Report')
            
            workbook = writer.book
            worksheet = writer.sheets['QA Report']
            
            # --- Define Professional Styles ---
            header_font = Font(bold=True, size=12, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1E293B") # Corporate Dark Slate
            wrap_alignment = Alignment(wrap_text=True, vertical='top')
            thin_border = Border(left=Side(style='thin', color="CCCCCC"), right=Side(style='thin', color="CCCCCC"), 
                                 top=Side(style='thin', color="CCCCCC"), bottom=Side(style='thin', color="CCCCCC"))
            
            # --- Apply Header Styles ---
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
            # --- Set Column Widths ---
            col_widths = { 'A': 25, 'B': 15, 'C': 40, 'D': 40, 'E': 45, 'F': 35, 'G': 35, 'H': 35, 'I': 15, 'J': 45 , 'K': 15 }
            for col_letter, width in col_widths.items():
                worksheet.column_dimensions[col_letter].width = width
                
            # --- Apply Data Cell Styles and Expand Row Heights for Images ---
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                # Only expand rows that have actual test data (skip the total row)
                if row_idx < worksheet.max_row:
                    worksheet.row_dimensions[row_idx].height = 160  # 🔴 Massive row height to fit the screenshot perfectly
                    
                for cell in row:
                    cell.alignment = wrap_alignment
                    cell.border = thin_border
                    
                    # Highlight Failures in Red
                    if cell.column == 9: # Pass/Fail Column
                        if cell.value == "FAIL":
                            cell.font = Font(bold=True, color="FF0000")
                        elif cell.value == "PASS":
                            cell.font = Font(bold=True, color="00B050")
            
            # --- Embed Actual Images into the Excel Cells ---
            # Columns 6 (F), 7 (G), 8 (H) map to Before, Intent, After
            for r_idx, row_data in enumerate(excel_rows, start=2):
                for col_idx, col_name in enumerate(["Before Image", "Intent Image", "After Image"], start=6):
                    img_path = row_data[col_name]
                    cell_ref = worksheet.cell(row=r_idx, column=col_idx)
                    cell_ref.value = "" # Erase the ugly text path!
                    
                    if img_path and os.path.exists(img_path):
                        try:
                            img = xlImage(img_path)
                            img.height = 190  # Scale image down to fit neatly in the cell
                            img.width = 250
                            worksheet.add_image(img, cell_ref.coordinate)
                        except Exception as e:
                            print(f"Warning: Failed to embed image {img_path}. Error: {e}")

            # --- Apply Grand Total Row Styles (Last Row) ---
            for cell in worksheet[worksheet.max_row]:
                cell.font = Font(bold=True, size=12)
                cell.fill = PatternFill("solid", fgColor="F1F5F9") # Light Gray
                
    # 🔴 EXPORT JSON TELEMETRY FOR THE UI TO READ
    run_data = {
        "total_chunks": total_chunks,
        "total_bugs": total_bugs,
        "details": excel_rows
    }
    with open(os.path.join(REPORT_DIR, "latest_run.json"), "w", encoding="utf-8") as f:
        json.dump(run_data, f)
                
    print("\n" + "="*70)
    print("🏁 SEMANTIC WORKFLOW COMPLETE")
    print(f"   Semantic Chunks Evaluated: {total_chunks}")
    print(f"   Confirmed Bugs Found: {total_bugs}")
    print(f"   Excel Report Saved To: {excel_report_path}")
    print(f"   💸 Total LLM Cost: ${total_cost:.4f}")
    print(f"   🪙 Total Tokens Used: {total_tokens}")
    print("="*70 + "\n")

if __name__ == "__main__":
    main()
